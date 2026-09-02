from pathlib import Path
import docx
from app.models.document import ParsedDocument, DocumentSection
from app.services.parsers.base import BaseDocumentParser, DocumentProcessingError


class DocxParser(BaseDocumentParser):
    """Parser for Microsoft Word (.docx) documents."""

    def parse(
        self,
        file_path: Path,
        filename: str,
        mime_type: str,
        attachment_id: str,
    ) -> ParsedDocument:
        file_size = file_path.stat().st_size if file_path.exists() else 0

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            err_msg = str(e)
            if str(file_path) in err_msg or str(file_path.parent) in err_msg:
                err_msg = "Corrupted or invalid DOCX archive."
            raise DocumentProcessingError(f"Corrupted or invalid DOCX document: {err_msg}")

        sections: list[DocumentSection] = []
        current_heading = None
        current_paragraphs: list[str] = []
        section_idx = 0

        def flush_section():
            nonlocal section_idx, current_heading, current_paragraphs
            if current_paragraphs or current_heading:
                body = "\n\n".join(current_paragraphs).strip()
                full_content = (
                    f"{current_heading}\n\n{body}"
                    if current_heading and body
                    else (current_heading or body)
                )
                if full_content.strip():
                    sections.append(
                        DocumentSection(
                            index=section_idx,
                            title=current_heading,
                            content=full_content,
                            page_number=1,
                            character_count=len(full_content),
                            metadata={"format": "docx", "heading": current_heading},
                        )
                    )
                    section_idx += 1
                current_paragraphs = []

        # Iterate through paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else "Normal"
            is_heading = style_name.startswith("Heading") or style_name in ["Title", "Subtitle"]

            if is_heading:
                flush_section()
                current_heading = text
            else:
                current_paragraphs.append(text)

        # Iterate through tables and append them
        for table_idx, table in enumerate(doc.tables):
            table_rows: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(row_cells):
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                flush_section()
                table_content = "\n".join(table_rows)
                sections.append(
                    DocumentSection(
                        index=section_idx,
                        title=f"Table {table_idx + 1}",
                        content=table_content,
                        page_number=1,
                        character_count=len(table_content),
                        metadata={"format": "docx", "table_index": table_idx},
                    )
                )
                section_idx += 1

        flush_section()

        if not sections:
            sections.append(
                DocumentSection(
                    index=0,
                    title=None,
                    content="",
                    page_number=1,
                    character_count=0,
                    metadata={"format": "docx"},
                )
            )

        total_chars = sum(s.character_count for s in sections)

        return ParsedDocument(
            attachment_id=attachment_id,
            filename=filename,
            mime_type=mime_type
            or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size=file_size,
            status="processed",
            total_characters=total_chars,
            page_count=1,
            section_count=len(sections),
            processed_at=self.get_current_timestamp(),
            sections=sections,
            metadata={"format": "docx", "section_count": len(sections)},
        )
